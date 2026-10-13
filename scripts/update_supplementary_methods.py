from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "outputs" / "submission_package" / "supplementary_methods.docx"


REPLACEMENTS = {
    "For the CNV sensitivity audit,": "For the CNV sensitivity audit,",
    "Externally defined BACH1 resources were evaluated": "Externally defined BACH1 resources were evaluated",
    "Directional mean-z scores were calculated": "Directional mean-z scores were calculated",
    "Patient-state units with at least ten cells were used": "Patient-state units with at least ten cells were used",
    "pySCENIC v0.12.1 was used": "pySCENIC v0.12.1 was used",
    "Tumour scATAC-seq context used": "Tumour scATAC-seq context used",
    "The 14 genes recovered in at least four of five pySCENIC seeds were classified": "The 14 genes recovered in at least four of five pySCENIC seeds were classified",
    "Spatial concordance was quantified": "Spatial concordance was quantified",
}


UPDATED = {
    "For the CNV sensitivity audit,": (
        "For the CNV sensitivity audit, pooled state fractions used all primary-set cells, whereas "
        "consensus-support proportions were also summarised at the patient-by-state level; units with "
        "fewer than 10 primary malignant epithelial cells were excluded from that descriptive display. "
        "For example, Stress-associated support was 1,831/3,780 (48.44%) in the pooled cell-level summary "
        "and 1,829/3,766 (48.57%) among eligible patient-state units; such small differences arise from "
        "the patient-state cell-count filter. Pooled cell-level proportions are weighted by the number of "
        "cells contributed by each patient, whereas patient-level summaries assign one observation to each "
        "eligible patient-state unit. Consequently, marked differences between the two summaries can arise "
        "from unequal patient contributions. In Supplementary Figure 1d, each point represents one eligible "
        "patient-state unit's consensus-support fraction (consensus cells divided by primary-set cells), "
        "and each black horizontal segment represents the median across eligible units within that state. "
        "States with fewer than four eligible patient-state units were interpreted descriptively only. "
        "The restricted set contained 4,906 cells from three patients, with 4,345 cells from AL08 (88.6%), "
        "507 from AL11 and 54 from P0031. The de-overlapped DoRothEA–hypoxia score association was "
        "recalculated after cell-level restriction and summarised by patient means as an exploratory "
        "descriptive sensitivity only; the restricted-set estimate (Spearman ρ = 0.500, n=3) was not "
        "treated as independent replication or used for inferential claims because of the restricted set's "
        "patient imbalance."
    ),
    "Directional mean-z scores were calculated": (
        "Directional mean-z scores were calculated from cell-level log-normalised expression and signed "
        "gene weights after filtering and duplicate collapse. For a cell i, gene set G and source weight "
        "w_g, the score was S_i = Σ(g∈G) w_g z_ig / Σ(g∈G) |w_g|, where z_ig is the gene-wise standardised "
        "expression value calculated across all 13,694 primary malignant epithelial cells. Source-signed "
        "weights were retained; zero-variance genes were assigned z=0, and the denominator was the sum of "
        "absolute weights for the genes present in the analysis matrix. Patient-level score summaries were "
        "arithmetic means of cell-level scores within patient, whereas patient-pseudobulk ssGSEA and "
        "decoupler ULM and MLM estimates were retained as sensitivity implementations. These alternative "
        "estimates were used to assess resource and scoring-method dependence rather than to define "
        "additional primary hypotheses."
    ),
    "Patient-state units with at least ten cells were used": (
        "Patient-state units with at least ten cells were used for within-state Spearman correlations. "
        "For the patient-level attenuation analysis, state proportions were calculated within patient, "
        "centred across patients and reduced by singular-value decomposition; the first two composition "
        "PCs were retained and z-standardised. The attenuation models adjusted for dataset, log1p-transformed "
        "median total counts, median detected-gene count, median mitochondrial percentage and the two "
        "composition PCs. Both scores were residualised with this same covariate design before the "
        "descriptive residual Spearman correlation was calculated. For the three states with at least four "
        "evaluable units, two-sided exact P values were obtained by enumerating all rank permutations when "
        "n≤9 and were Holm-adjusted across states. Fisher-z 95% confidence intervals are approximate and "
        "were not used for inference in these small analyses. Given the 16 independent patients and the "
        "estimation of multiple nuisance covariates, residual correlations and their associated P values "
        "and intervals were interpreted as exploratory measures of attenuation, not as causal decomposition "
        "or independent validation. A CLR sensitivity used a 0.5-cell pseudocount before PCA."
    ),
    "pySCENIC v0.12.1 was used": (
        "pySCENIC v0.12.1 was used for exploratory regulatory-candidate context. Final recurrence analysis "
        "used the 4,906-cell restricted CNV-consensus set and the full transcription-factor search space. "
        "Five runs were completed with seeds 777–781. Each run used GRNBoost2, cisTarget motif pruning and "
        "AUCell scoring with hg38 10-kb and 500-bp cisTarget v10 ranking databases and motif-v10nr "
        "annotations. BACH1 was excluded from target-frequency tables to avoid self-target inflation. "
        "Seed-specific outputs were retained as diagnostics; cross-seed recurrence was used for candidate "
        "prioritisation. Cross-seed recurrence reflects computational stability under repeated network "
        "inference rather than biological reproducibility across independent patients. Because the "
        "restricted CNV-consensus set contained only three patients and was strongly dominated by AL08, "
        "these results were used exclusively for exploratory candidate prioritisation."
    ),
    "Tumour scATAC-seq context used": (
        "Tumour scATAC-seq context used author-filtered peak-by-barcode HDF5 matrices and quality-control "
        "summaries for AL05, AL08, AL10, AL11 and AL12, comprising 28,500 cells and 640,447 sample-specific "
        "peaks. FASTQ files were not reprocessed, peak calling was not repeated and no additional hard "
        "cell-level filtering was imposed."
    ),
    "BACH1 (JASPAR MA1633.2)": (
        "BACH1 (JASPAR MA1633.2) and BACH1::MafK (MA0591.2) motif coordinates were obtained from the "
        "UCSC hg38 JASPAR2026 bigBed resource accessed on 10–11 September 2026. Merged GENCODE release 44 "
        "TSS ±10 kb windows for recurrent candidates were queried with bigBedToBed and intersected with the "
        "five author-filtered scATAC peak matrices. Records retained motif model, UCSC score, sample identity, "
        "accessibility and the non-negative genomic gap between the nearest peak boundary and the annotated TSS; "
        "the distance was set to zero when the peak interval overlapped the TSS and otherwise was the gap to the "
        "nearest boundary, without strand adjustment. Window membership was defined by interval overlap with the "
        "TSS window. Promoter-window support was defined as a motif-overlapping accessible peak within ±2 kb of "
        "the annotated TSS. Score thresholds of 300, 400 and 500 were summarised descriptively; the integrated "
        "candidate matrix used the threshold of ≥300. At the unfiltered link level, the analysis comprised 497 "
        "unique sample-specific motif-positive peaks and 704 motif–peak–candidate proximity links, including 464 "
        "BACH1 links and 240 BACH1::MafK links. These unfiltered records are referred to as unfiltered "
        "proximity-link context. The motif-model, candidate-ranking and example-locus panels of Supplementary "
        "Figure 2 display all retained motif-overlapping links before score-threshold filtering, whereas the "
        "threshold-sensitivity panel reports thresholded proximity-link evidence and Supplementary Figure 3b uses "
        "the ≥300 threshold for candidate tiering. The number of unique motif-positive peaks differs from the "
        "number of motif–peak–candidate proximity links because the latter counts retained motif-model–peak–candidate "
        "records. The same accessible peak may therefore contribute separate records for the two motif models; "
        "207 of the 497 unique sample-specific peaks had both motif models, whereas no peak was assigned to more "
        "than one candidate gene."
    ),
    "The 14 genes recovered in at least four of five pySCENIC seeds were classified": (
        "The 14 genes recovered in at least four of five pySCENIC seeds were classified using recurrence and "
        "chromatin context. Tier 1 required recovery in all five seeds, motif-overlapping accessible peaks "
        "in all five scATAC samples and promoter-window support at a UCSC motif score of ≥300. Tier 2 "
        "required recovery in exactly four seeds, five-sample motif context at the same threshold and "
        "promoter-window support. Remaining recurrent genes were labelled context-limited. GRNBoost2 "
        "importance was summarised as the recurrence-weighted mean importance across runs; partial BACH1 "
        "transcript correlation was the Pearson correlation between candidate-gene expression and BACH1 mRNA "
        "expression after both were residualised for dataset, sample identity, epithelial Leiden cluster, "
        "total counts, detected genes and mitochondrial percentage; and stress-state enrichment was the mean log-normalised "
        "expression in Stress-associated cells minus the mean across the other primary malignant epithelial "
        "states. The partial correlation did not use the DoRothEA score. These matrix columns, together with motif-link counts, were shown for context only and did "
        "not enter tier assignment. The classes are evidence-prioritisation categories, not confidence "
        "intervals or claims of direct regulation."
    ),
    "Spatial concordance was quantified": (
        "Spatial concordance was quantified with bivariate Moran's I using a six-neighbour graph constructed "
        "from array coordinates; edge weights were row-standardised equally across each spot's available "
        "neighbours. For Figure 5d, raw DoRothEA BACH1 and raw Hallmark hypoxia scores were compared with a "
        "20-by-20 locally constrained block null within each section. The DoRothEA score was held fixed and "
        "the hypoxia score labels were permuted within blocks. For Supplementary Figure 5, the same procedure "
        "was applied to the state-PC-residualised score pair at block sizes of 10, 15, 20 and 30 array "
        "coordinates. Blocks were defined by integer array-row and array-column coordinates divided by the "
        "corresponding block size; blocks containing one spot were unchanged. For each analysis, 1,000 "
        "permutations were performed and the upper-tail empirical P value was calculated as (b+1)/(B+1), "
        "where b is the number of null Moran's I values at least as large as observed. The 95% permutation "
        "intervals shown in the figures are the 2.5th and 97.5th percentiles of the corresponding null "
        "distribution. These block-restricted nulls preserve coarse local block structure and the within-block "
        "score distribution but not the full variogram. Spatial summaries are section-level and exploratory; "
        "the single LUSC section was interpreted descriptively."
    ),
}


def update_document() -> None:
    document = Document(DOCX)
    replaced = set()
    for paragraph in document.paragraphs:
        text = paragraph.text
        for prefix, replacement in UPDATED.items():
            if text.startswith(prefix):
                paragraph.text = replacement
                replaced.add(prefix)
                break

    missing = set(UPDATED) - replaced
    if missing:
        raise RuntimeError(f"Could not locate paragraphs: {sorted(missing)}")

    for paragraph in document.paragraphs:
        if paragraph.text:
            paragraph.text = paragraph.text.replace("BACH1::Mafk", "BACH1::MafK")
            paragraph.text = paragraph.text.replace("TSS±10-kb", "TSS ±10 kb")
            paragraph.text = paragraph.text.replace("Epithelial_core", "Epithelial core")
            paragraph.text = paragraph.text.replace("AT2_like", "AT2-like")

    document.save(DOCX)


if __name__ == "__main__":
    update_document()
