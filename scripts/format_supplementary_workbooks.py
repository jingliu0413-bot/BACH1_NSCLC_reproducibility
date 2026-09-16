#!/usr/bin/env python3
"""Format supplementary workbook tabs for submission readability."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "manuscript_submission_20260914"
WORKBOOKS = [OUT / "Supplementary_Tables.xlsx", OUT / "Supplementary Tables.xlsx"]

GROUP_FILLS = {
    "S1": "D9EAF7",
    "S2": "E7F3EA",
    "S3": "F7E9C8",
    "S4": "F8E9E7",
    "S5": "ECE7F2",
}


def width_for(values: list[str], header: str) -> float:
    max_len = max([len(header)] + [len(v) for v in values[:200] if v is not None])
    if max_len <= 8:
        return 10
    if max_len <= 14:
        return 14
    if max_len <= 24:
        return 20
    if max_len <= 45:
        return 30
    return 42


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        prefix = ws.title[:2]
        fill = PatternFill("solid", fgColor=GROUP_FILLS.get(prefix, "D9EAF7"))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float):
                    header = str(ws.cell(row=1, column=cell.column).value or "").lower()
                    if any(k in header for k in ["rho", "beta", "hazard", "hr", "correlation", "partial", "jaccard"]):
                        cell.number_format = "0.000"
                    elif any(k in header for k in ["pct", "percent", "fraction", "proportion"]):
                        cell.number_format = "0.0"
                    elif abs(cell.value) < 0.001 and cell.value != 0:
                        cell.number_format = "0.00E+00"
                    else:
                        cell.number_format = "0.000"

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header = str(ws.cell(row=1, column=col_idx).value or "")
            values = [str(ws.cell(row=row_idx, column=col_idx).value or "") for row_idx in range(2, min(ws.max_row, 201) + 1)]
            width = width_for(values, header)
            if header.lower() in {"gene", "tf", "state", "cohort", "comparison", "candidate_tier", "resource_or_package"}:
                width = max(width, 18)
            if any(k in header.lower() for k in ["genes", "samples", "details", "notes", "source", "members"]):
                width = max(width, 34)
            ws.column_dimensions[col_letter].width = width

        ws.row_dimensions[1].height = 34

    wb.save(path)
    print(f"Formatted {path}")


def main() -> None:
    for workbook in WORKBOOKS:
        format_workbook(workbook)


if __name__ == "__main__":
    main()
