"""Remove orphaned supplementary clinical/NOD-adjacent outputs from final package."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/manuscript_submission_20260914"

OBSOLETE_SHEETS = {"S5_state_signature", "S5_TCGA_clinical", "S5_TCGA_OS_Cox"}
OBSOLETE_FILE_STEMS = [
    OUT / "figures/Figure5_multiomic_candidate_working_model",
    OUT / "figures/Supplementary_Figure3_current_candidate_integration",
    OUT / "figures/Supplementary_Figure4_TCGA_clinical_context_projection",
]
OBSOLETE_FILES = [
    OUT / "tables/stress_ap1_state_signature_genes.csv",
    OUT / "tables/tcga_stress_ap1_state_signature_clinical_associations.csv",
    OUT / "tables/tcga_stress_ap1_state_signature_os_cox.csv",
    OUT / "tables/tcga_stress_ap1_state_signature_scores_clinical.csv",
    OUT / "reports/state_contextual_clinical_revision_summary.json",
]


def remove_workbook_sheets(path: Path) -> None:
    if not path.exists():
        return
    wb = load_workbook(path)
    changed = False
    for sheet in list(wb.sheetnames):
        if sheet in OBSOLETE_SHEETS:
            del wb[sheet]
            changed = True
    if changed:
        wb.save(path)


def filter_flat_csv(path: Path) -> None:
    if not path.exists():
        return
    tmp = path.with_suffix(path.suffix + ".tmp")
    with path.open("r", newline="", encoding="utf-8") as fh, tmp.open("w", newline="", encoding="utf-8") as out:
        reader = csv.DictReader(fh)
        writer = csv.DictWriter(out, fieldnames=reader.fieldnames)
        writer.writeheader()
        for row in reader:
            if row.get("source_sheet") in OBSOLETE_SHEETS:
                continue
            writer.writerow(row)
    tmp.replace(path)


def main() -> None:
    for workbook in [OUT / "Supplementary_Tables.xlsx", OUT / "Supplementary Tables.xlsx"]:
        remove_workbook_sheets(workbook)

    for flat in [OUT / "Supplementary_Tables.csv", OUT / "Supplementary_Tables_revision_source_data.csv"]:
        filter_flat_csv(flat)

    for stem in OBSOLETE_FILE_STEMS:
        for ext in [".pdf", ".png", ".svg"]:
            path = stem.with_suffix(ext)
            if path.exists():
                path.unlink()

    for path in OBSOLETE_FILES:
        if path.exists():
            path.unlink()


if __name__ == "__main__":
    main()
