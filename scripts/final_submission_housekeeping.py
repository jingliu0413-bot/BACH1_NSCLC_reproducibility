#!/usr/bin/env python3
"""Final submission housekeeping for labels and supplementary workbook tabs."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "manuscript_submission_20260914"

TIER_REPLACEMENTS = {
    "Tier 1: 5/5 seeds + 5-sample ATAC": "Tier 1: 5/5 seeds + 5-sample ATAC + promoter support",
    "Tier 2: 4/5 seeds + promoter ATAC": "Tier 2: 4/5 seeds + 5-sample ATAC + promoter support",
}

SHEET_RENAMES = {
    "S4_state_summary": "S4E_state_summary",
    "S4_attenuation": "S4F_attenuation",
    "S4_within_state": "S4G_within_state",
    "S4_candidate_matrix": "S4H_candidate_matrix",
    "S4_state_malignancy_audit": "S4I_state_malignancy_audit",
    "S4_IE_sensitivity_state": "S4J_IE_sensitivity_state",
    "S4_IE_sensitivity_genes": "S4K_IE_sensitivity_genes",
    "S4_IE_overlap": "S4L_IE_overlap",
    "S4_CLR_composition_sens": "S4M_CLR_composition_sens",
    "S5_TF_all_benchmark": "S5A_TF_all_benchmark",
    "S5_TF_selected": "S5B_TF_selected",
    "S5_scRNA_split": "S5C_scRNA_split",
}

DESIRED_SHEET_ORDER = [
    "S1A_signature_genes",
    "S1B_score_summary",
    "S1C_overlap_summary",
    "S1D_overlap_genes",
    "S2A_null_parameters",
    "S2B_null_summary",
    "S2C_input_gene_counts",
    "S2D_first50_matches",
    "S3A_atac_summary",
    "S3B_target_support",
    "S3C_score_thresholds",
    "S3D_missing_targets",
    "S4A_pyscenic_runs",
    "S4B_pyscenic_jaccard",
    "S4C_recurrent_targets",
    "S4D_resources",
    "S4E_state_summary",
    "S4F_attenuation",
    "S4G_within_state",
    "S4H_candidate_matrix",
    "S4I_state_malignancy_audit",
    "S4J_IE_sensitivity_state",
    "S4K_IE_sensitivity_genes",
    "S4L_IE_overlap",
    "S4M_CLR_composition_sens",
    "S5A_TF_all_benchmark",
    "S5B_TF_selected",
    "S5C_scRNA_split",
]


def replace_text_file(path: Path) -> None:
    if not path.exists():
        return
    text = path.read_text()
    for old, new in {**TIER_REPLACEMENTS, **SHEET_RENAMES}.items():
        text = text.replace(old, new)
    path.write_text(text)


def update_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for old, new in SHEET_RENAMES.items():
        if old in wb.sheetnames:
            wb[old].title = new

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    value = cell.value
                    for old, new in TIER_REPLACEMENTS.items():
                        value = value.replace(old, new)
                    if value != cell.value:
                        cell.value = value

    ordered = [wb[name] for name in DESIRED_SHEET_ORDER if name in wb.sheetnames]
    ordered += [ws for ws in wb.worksheets if ws.title not in DESIRED_SHEET_ORDER]
    wb._sheets = ordered
    wb.save(path)


def main() -> None:
    for rel in [
        "tables/figure4e_candidate_evidence_matrix.csv",
        "tables/figure5_candidate_tier_matrix.csv",
        "Supplementary_Tables.csv",
        "Supplementary_Tables_revision_source_data.csv",
    ]:
        replace_text_file(OUT / rel)
    update_workbook(OUT / "Supplementary_Tables.xlsx")


if __name__ == "__main__":
    main()
